"""Validate an official Seoul 121-place Excel list without guessing coordinates.

Usage: python -m scripts.validate_places /path/to/official.xlsx
Official file: https://data.seoul.go.kr/dataList/OA-21285/F/1/datasetView.do
Does not mutate the production catalog. Official polygons/centroids must be separately verified.
"""
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

def validate(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    candidates=[]
    for ws in wb.worksheets:
        rows=list(ws.iter_rows(values_only=True))
        for i,row in enumerate(rows):
            labels=[str(v).strip() if v is not None else "" for v in row]
            if any(v in ("AREA_NM","장소명","장소이름") for v in labels):
                name_col=next(j for j,v in enumerate(labels) if v in ("AREA_NM","장소명","장소이름"))
                code_col=next((j for j,v in enumerate(labels) if v in ("AREA_CD","장소코드","장소 코드")),None)
                if code_col is None:continue
                records=[]
                for data in rows[i+1:]:
                    n=str(data[name_col]).strip() if len(data)>name_col and data[name_col] is not None else ""
                    c=str(data[code_col]).strip() if len(data)>code_col and data[code_col] is not None else ""
                    if n and c:records.append({"name":n,"code":c})
                if len(records)>len(candidates):candidates=records
    codes=[x["code"] for x in candidates];names=[x["name"] for x in candidates]
    if len(candidates)!=121 or len(set(codes))!=121 or len(set(names))!=121:
        raise ValueError(f"Expected 121 unique names and codes; got {len(candidates)}. Check official workbook columns.")
    return candidates

if __name__=="__main__":
    data=validate(Path(sys.argv[1]))
    print(json.dumps({"verified_count":len(data),"items":data},ensure_ascii=False,indent=2))
